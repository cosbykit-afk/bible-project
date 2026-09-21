#!/usr/bin/env python3
"""Canonical ingest of Kit's Bible files + diff vs Prophets.xlsx.
Outputs staging CSVs + diff report to /home/hatch/workspace/bible-project/worker-out/canon/.
"""
import openpyxl, csv, re, json, os
from collections import Counter

OUT = '/home/hatch/workspace/bible-project/worker-out/canon'
os.makedirs(OUT, exist_ok=True)
BIBLE = '/home/hatch/workspace/user/files/7617FAE6-744D-4ABA-A3B9-53CCC19804E6-BibleFull (v 1).xlsx'
PROPH = '/tmp/bible/Prophets.xlsx'
HEB = re.compile(r'[\u05d0-\u05ea]')

def letters(s):
    return ''.join(HEB.findall(s or ''))

log = []
def say(m):
    log.append(m); print(m, flush=True)

say('=== Loading BibleFull Bible sheet ===')
wb = openpyxl.load_workbook(BIBLE, read_only=True, data_only=True)
ws = wb['Bible']
rows = []
n = 0
for r in ws.iter_rows(min_row=2, values_only=True):
    n += 1
    rows.append(r)
say(f'rows: {n}')
ids = [r[0] for r in rows]
say(f'id continuity: min={min(ids)} max={max(ids)} unique={len(set(ids))} sequential={sorted(ids)==list(range(1,n+1))}')

# word_pos within verse (file order)
pos_counter = {}
for r in rows:
    key = (r[2], r[3], r[4])
    pos_counter[key] = pos_counter.get(key, 0) + 1

pos2 = {}
canon = []
for r in rows:
    key = (r[2], r[3], r[4])
    pos2[key] = pos2.get(key, 0) + 1
    canon.append({
        'id': r[0], 'book': r[2], 'chapter': r[3], 'verse': r[4],
        'word_pos': pos2[key],
        'word_pointed': r[5] or '', 'no_vowel_raw': r[6] or '',
        'letters': letters(r[6]),
        'prefix1': r[7] or '', 'prefix2': r[8] or '', 'prefix3': r[9] or '',
        'base_word': r[10] or '', 'suffix1': r[11] or '', 'suffix2': r[12] or '',
        'strongs': str(r[13]).strip() if r[13] else '',
        'base_def': r[14] or '', 'translation': r[15] or '',
        'proper_noun': r[16] or '',
    })

with open(f'{OUT}/canon_words.csv', 'w', encoding='utf-8', newline='') as f:
    w = csv.DictWriter(f, fieldnames=list(canon[0].keys()))
    w.writeheader(); w.writerows(canon)
say(f'canon_words.csv written: {len(canon)} rows')

# Books sheet
wsb = wb['Books']
books = []
for r in wsb.iter_rows(min_row=2, values_only=True):
    if r[0] is None: continue
    books.append({'book_num': int(r[0]), 'name_en': (r[1] or '').strip(), 'name_he': (r[2] or '').strip()})
with open(f'{OUT}/canon_books.csv', 'w', encoding='utf-8', newline='') as f:
    w = csv.DictWriter(f, fieldnames=['book_num', 'name_en', 'name_he'])
    w.writeheader(); w.writerows(books)
say(f'books: {len(books)}')

# per-book word counts
bc = Counter(c['book'] for c in canon)
say('per-book counts: ' + json.dumps({b['book_num']: bc[b['book_num']] for b in books}))

# prefilled strongs/translation stats
say(f"prefilled strongs: {sum(1 for c in canon if c['strongs'])}, translations: {sum(1 for c in canon if c['translation'])}")

# === Sheet1 notes: verify col A joins to word id; Aramaic flags ===
say('=== Sheet1 notes ===')
ws1 = wb['Sheet1']
notes_total = 0; a_notes = 0; join_ok = 0; join_bad = 0; bad_samples = []
aramaic_ids = set()
idmap = {c['id']: c for c in canon}
for r in ws1.iter_rows(min_row=2, values_only=True):
    notes_total += 1
    nid, nb, nc, nv, nword = r[0], r[2], r[3], r[4], r[6]
    if r[7] == 'A':
        a_notes += 1
        aramaic_ids.add(nid)
    # join check on a sample
    if notes_total % 5000 == 0 or notes_total < 10:
        c = idmap.get(nid)
        if c and (c['book'], c['chapter'], c['verse']) == (nb, nc, nv) and letters(nword) == c['letters']:
            join_ok += 1
        else:
            join_bad += 1
            if len(bad_samples) < 5: bad_samples.append((nid, nb, nc, nv, nword))
say(f'notes rows: {notes_total}, Language=A: {a_notes}, join check ok={join_ok} bad={join_bad} samples={bad_samples}')
with open(f'{OUT}/aramaic_word_ids.txt', 'w') as f:
    f.write('\n'.join(str(i) for i in sorted(aramaic_ids)))
say(f'aramaic_word_ids.txt: {len(aramaic_ids)} ids')

# === Diff Bible sheet vs Prophets.xlsx (books 6-39) ===
say('=== Diff vs Prophets.xlsx ===')
wp = openpyxl.load_workbook(PROPH, read_only=True, data_only=True)
say(f'Prophets sheets: {wp.sheetnames}')
junk_log = []
canon_key = {}
for c in canon:
    if c['book'] >= 6:
        canon_key[(c['book'], c['chapter'], c['verse'], c['word_pos'])] = c
total_p = 0; match = 0; mism = []; sheet_rows = {}; mismatch_verses = Counter()
for name in wp.sheetnames:
    wsh = wp[name]
    sheet_rows[name] = (wsh.max_row - 1, wsh.max_column)
for name in wp.sheetnames:
    wsh = wp[name]
    nr, nc = sheet_rows[name]
    if nc > 13:
        junk_log.append(f'{name}: {nc} columns (expected <=13); extra columns ignored as junk')
    plist = []
    for r in wsh.iter_rows(min_row=2, values_only=True):
        plist.append((r[2], r[3], r[4], r[5] or '', r[6] or ''))
    ppos = {}
    for (b, ch, v, pointed, novow) in plist:
        key = (b, ch, v)
        ppos[key] = ppos.get(key, 0) + 1
        total_p += 1
        c = canon_key.get((b, ch, v, ppos[key]))
        if c is None:
            mism.append(('NO_CANON', name, b, ch, v, ppos[key]))
            continue
        if c['letters'] == letters(novow) and (c['word_pointed'] or '') == (pointed or ''):
            match += 1
        else:
            if len(mism) < 30:
                mism.append(('DIFF', name, b, ch, v, ppos[key], c['word_pointed'], pointed, c['letters'], letters(novow)))
            mismatch_verses[(b, ch, v)] += 1

say(f'prophets word rows compared: {total_p}, exact matches: {match} ({100*match/max(total_p,1):.3f}%)')
say(f'mismatches recorded: {len(mism)}; verses affected: {len(mismatch_verses)}')
for m in mism[:15]:
    say('  ' + str(m)[:200])
for j in junk_log:
    say('JUNK: ' + j)
say('sheet rows/cols: ' + json.dumps(sheet_rows))

# canonical coverage: prophets sheets cover which books?
pb = Counter()
for name in wp.sheetnames:
    wsh = wp[name]
    for r in wsh.iter_rows(min_row=2, max_col=3, values_only=True):
        if r[2] is not None: pb[r[2]] += 1
        break
say('prophets books present (sampled): ' + json.dumps(dict(pb)))

report = {
    'canon_rows': len(canon),
    'id_sequential': sorted(ids) == list(range(1, n + 1)),
    'prefilled_strongs': sum(1 for c in canon if c['strongs']),
    'prefilled_translations': sum(1 for c in canon if c['translation']),
    'notes_total': notes_total, 'aramaic_notes': a_notes,
    'notes_join_ok': join_ok, 'notes_join_bad': join_bad,
    'prophets_rows': total_p, 'prophets_exact_match': match,
    'prophets_match_pct': 100 * match / max(total_p, 1),
    'mismatch_count': len(mism), 'mismatch_verses': len(mismatch_verses),
    'junk_columns': junk_log,
    'per_book': {b['book_num']: bc[b['book_num']] for b in books},
}
with open(f'{OUT}/canon_report.json', 'w') as f:
    json.dump(report, f, indent=2, ensure_ascii=False)
with open(f'{OUT}/canon_diff.log', 'w') as f:
    f.write('\n'.join(log))
say('DONE')
