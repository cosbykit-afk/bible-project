#!/usr/bin/env python3
"""Map Sheet1 notes rows to canonical words via orig id; derive Aramaic flags on surrogate ids."""
import openpyxl, csv, json, re
from collections import Counter, defaultdict

OUT = '/home/hatch/workspace/bible-project/worker-out/canon'
HEB = re.compile(r'[\u05d0-\u05ea]')
def letters(s): return ''.join(HEB.findall(s or ''))

canon = []
with open(f'{OUT}/canon_words.csv', encoding='utf-8') as f:
    for i, r in enumerate(csv.DictReader(f), start=1):
        r['word_id'] = i
        canon.append(r)
by_orig = defaultdict(list)
for c in canon:
    by_orig[str(c['id'])].append(c)

wb = openpyxl.load_workbook('/home/hatch/workspace/user/files/7617FAE6-744D-4ABA-A3B9-53CCC19804E6-BibleFull (v 1).xlsx',
                            read_only=True, data_only=True)
ws = wb['Sheet1']
mapped = 0; unmapped = 0; dupe_resolved = 0; dupe_ambiguous = 0
unmapped_samples = []
aram_ids = set()
note_id_samples = []
n = 0
for r in ws.iter_rows(min_row=2, values_only=True):
    n += 1
    nid = str(r[0]); nb, nc, nv = r[2], r[3], r[4]
    nletters = letters(r[6]); lang = r[7]
    cands = by_orig.get(nid, [])
    chosen = None
    if len(cands) == 1:
        chosen = cands[0]; mapped += 1
    elif len(cands) > 1:
        # disambiguate by ref+letters
        for cd in cands:
            if (int(cd['book']), int(cd['chapter']), int(cd['verse'])) == (nb, nc, nv) and cd['letters'] == nletters:
                chosen = cd; dupe_resolved += 1; break
        if chosen is None:
            dupe_ambiguous += 1
            if len(unmapped_samples) < 5:
                unmapped_samples.append(('DUPE_AMBIG', nid, nb, nc, nv, nletters))
    else:
        unmapped += 1
        if len(unmapped_samples) < 10:
            unmapped_samples.append(('NO_ORIG', nid, nb, nc, nv, nletters, str(r[6])[:40]))
    if lang == 'A' and chosen is not None:
        aram_ids.add(chosen['word_id'])
    if n <= 5:
        note_id_samples.append((nid, r[1], nb, nc, nv))

print(f'notes rows: {n}, mapped: {mapped}, dupe_resolved: {dupe_resolved}, dupe_ambiguous: {dupe_ambiguous}, unmapped: {unmapped}')
print('unmapped samples:', unmapped_samples)
print('first rows (id, Notes-col, book,ch,v):', note_id_samples)
print('aramaic canon word_ids:', len(aram_ids))

with open(f'{OUT}/aramaic_word_ids.txt', 'w') as f:
    f.write('\n'.join(str(i) for i in sorted(aram_ids)))

# what is the notes ordering? check col B ('Notes') content distribution
notes_col = Counter()
for r in ws.iter_rows(min_row=2, max_col=2, values_only=True):
    notes_col[str(r[1])[:30] if r[1] else ''] += 1
    if sum(notes_col.values()) > 200000: break
print('Notes-col sample values:', notes_col.most_common(8))

report = {'notes_rows': n, 'mapped': mapped, 'dupe_resolved': dupe_resolved,
          'dupe_ambiguous': dupe_ambiguous, 'unmapped': unmapped,
          'unmapped_samples': [list(s) for s in unmapped_samples],
          'aramaic_word_count': len(aram_ids)}
with open(f'{OUT}/notes_mapping_report.json', 'w') as f:
    json.dump(report, f, indent=2, ensure_ascii=False)
print('DONE')
